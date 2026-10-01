import os
import glob
import shutil
import openpyxl

# Look for files in C:\$Recycle.Bin
base = 'C:\\$Recycle.Bin'
found = []
for root, dirs, files in os.walk(base):
    for f in files:
        if 'R7U1LCP' in f or 'RY9GHLF' in f or 'R9NJMZW' in f or 'R83JA0P' in f:
            src = os.path.join(root, f)
            dst = os.path.join(os.getcwd(), 'recovered_' + f + '.xlsm')
            shutil.copy2(src, dst)
            print(f'Copied {src} -> {dst}')
            found.append(dst)

for fpath in found:
    print('=== Analyzing:', fpath)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    print('Sheets:', wb.sheetnames)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f' Sheet {sname}: max_row={ws.max_row}, max_col={ws.max_column}')
        if sname == 'CONTABILIDAD':
            for r in range(8, min(14, ws.max_row+1)):
                print(f'   Row {r}:', [ws.cell(r, c).value for c in range(4, 15)])
        elif sname == 'ImportCONTABILIDAD':
            for r in range(1, min(6, ws.max_row+1)):
                print(f'   Row {r}:', [ws.cell(r, c).value for c in range(1, 12)])
